"""
Renderers and checks for the Publish step. No LLM calls in this module.

  render_docx   writes the manuscript INTO the user's own Word template: the
                template's sample text is removed, its page setup, columns,
                headers/footers and section breaks are kept, and every
                paragraph is given one of the template's own styles.
  render_latex  returns a .zip (the user's template files + manuscript.tex +
                references.bib + figures/) to compile in Overleaf or locally.
  check         the compliance checklist shown beside the download.
"""
from __future__ import annotations

import io
import re
import zipfile

from core.manuscript_template import _decode, as_docx_bytes, docx_sections, unpack_latex

CITE_RE = re.compile(r"\[(\d{1,3}(?:\s*(?:,|-|–)\s*\d{1,3})*)\]")
TODO_RE = re.compile(r"\[TO ADD:[^\]]*\]", re.I)

AI_DISCLOSURE = (
    "During the preparation of this work the authors used the Sift and Infinity research "
    "agents (Orcus Intelligence Lab) to search and summarise the literature and to draft "
    "parts of the text. The authors reviewed and edited the content as needed and take "
    "full responsibility for the content of the publication."
)


# ── Shared helpers ──────────────────────────────────────────────────────────

def cite_numbers(group: str) -> list[int]:
    out = []
    for part in re.split(r"\s*,\s*", group.strip()):
        m = re.fullmatch(r"(\d+)\s*[-–]\s*(\d+)", part)
        if m and int(m.group(1)) <= int(m.group(2)) <= int(m.group(1)) + 40:
            out.extend(range(int(m.group(1)), int(m.group(2)) + 1))
        elif part.strip().isdigit():
            out.append(int(part))
    return out


def blocks(text: str) -> list[tuple[str, str]]:
    """Section text -> [("h2", heading) | ("p", paragraph)]."""
    out, buf = [], []

    def flush():
        if buf:
            out.append(("p", " ".join(" ".join(buf).split())))
            buf.clear()

    for line in (text or "").splitlines():
        line = line.strip()
        m = re.match(r"^#{1,4}\s+(.+)$", line)
        if m:
            flush()
            out.append(("h2", m.group(1).strip().strip("*").strip()))
        elif not line:
            flush()
        else:
            buf.append(line)
    flush()
    return out


def _norm_heading(t: str) -> str:
    t = re.sub(r"^((\d+|[IVXLC]+)[.)]?\s+)?", "", (t or "").strip().strip("#*_ ").strip())
    return re.sub(r"[^a-z0-9]+", " ", t.lower().strip("*_ ")).strip()


def strip_own_heading(text: str, heading: str) -> str:
    """Models often open a section by repeating its heading ("## Introduction").
    The renderer prints the heading itself, so drop that first line."""
    lines = (text or "").lstrip().split("\n")
    if lines and len(lines[0]) < 160 and _norm_heading(lines[0]) == _norm_heading(heading):
        return "\n".join(lines[1:]).lstrip()
    return text or ""


def parse_authors(text: str) -> list[dict]:
    out = []
    for line in (text or "").splitlines():
        parts = [p.strip() for p in line.split(";")]
        if len(parts) == 1 and "," in line:         # typed with commas: "Name, Affiliation, email"
            bits = [b.strip() for b in line.split(",") if b.strip()]
            mail = next((b for b in bits if "@" in b), "")
            rest = [b for b in bits if b != mail]
            parts = [rest[0] if rest else "", ", ".join(rest[1:]), mail]
        if parts and parts[0]:
            out.append({"name": parts[0],
                        "affiliation": parts[1] if len(parts) > 1 else "",
                        "email": parts[2] if len(parts) > 2 else ""})
    return out[:40]


def _surname(authors: str) -> tuple[str, int]:
    names = [a.strip() for a in re.split(r",| and |;|&", authors or "") if a.strip()]
    names = [n for n in names if not re.fullmatch(r"et al\.?", n, re.I)]
    if not names:
        return "Anon", 0
    toks = names[0].replace(".", ". ").split()
    real = [t for t in toks if len(t.strip(".")) > 1]
    return (real[-1] if real else toks[-1]).strip(",."), len(names)


def author_year(p: dict) -> str:
    s, n = _surname(p.get("authors") or "")
    etal = " et al." if n > 2 or re.search(r"et al", p.get("authors") or "", re.I) else ""
    return f"{s}{etal}, {p.get('year') or 'n.d.'}"


REFERENCE_STYLES = ("ieee", "apa", "harvard", "vancouver", "acm")


def _people(authors: str) -> tuple[list[tuple[str, str]], bool]:
    """'Jane Q. Smith, A. Doe et al.' -> ([("Jane Q.", "Smith"), ("A.", "Doe")], True)."""
    raw = [a.strip() for a in re.split(r",|;| and |&", authors or "") if a.strip()]
    etal = any(re.search(r"et al\.?$", a, re.I) for a in raw)
    out = []
    for a in raw:
        a = re.sub(r"\s*et al\.?$", "", a, flags=re.I).strip()
        toks = a.replace(".", ". ").split()
        if not toks:
            continue
        if len(toks) > 1 and len(toks[-1].strip(".")) == 1 and len(toks[0].strip(".")) > 1:
            toks = toks[1:] + toks[:1]                  # "Smith J." -> given last; move family to the end
        out.append((" ".join(toks[:-1]), toks[-1].strip(",.")))
    return out, etal


def _initials(given: str, sep: str = " ", dot: str = ".") -> str:
    return sep.join(t[0].upper() + dot for t in given.replace("-", " ").split() if t[:1].isalpha())


def format_reference(p: dict, style: str, n: int | None, numbered: bool = True) -> str:
    """One reference-list entry in the venue's style, from the fields Sift
    has (authors, title, venue, year, link). Volume, issue and pages are not
    collected by the search, so they are not printed in any style."""
    people, etal = _people(p.get("authors") or "")
    title, venue = (p.get("title") or "").strip().rstrip("."), (p.get("venue") or "").strip()
    year, url = p.get("year") or "n.d.", (p.get("url") or "").strip()

    def join(names, last=" and ", oxford=True):
        if etal:
            return (names[0] if names else "Anon") + " et al."
        if len(names) <= 1:
            return names[0] if names else "Anon"
        if len(names) == 2:
            return names[0] + last + names[1]
        return ", ".join(names[:-1]) + ("," if oxford else "") + last + names[-1]

    num = f"[{n}] " if (n is not None and numbered) else ""
    if style == "vancouver":
        who = join([f"{f} {_initials(g, '', '')}".strip() for g, f in people], last=", ", oxford=False)
        return (f"{n}. " if (n is not None and numbered) else "") + f"{who}. {title}. " \
            + (f"{venue}. " if venue else "") + f"{year}." + (f" Available from: {url}" if url else "")
    if style in ("apa", "harvard"):
        who = join([f"{f}, {_initials(g)}".strip(", ") for g, f in people], last=", & " if style == "apa" else " and ",
                   oxford=False)
        return num + f"{who} ({year}). {title}. " + (f"{venue}. " if venue else "") + (url if url else "")
    if style == "acm":
        who = join([f"{g} {f}".strip() for g, f in people])
        return num + f"{who}. {year}. {title}. " + (f"{venue}. " if venue else "") + (url if url else "")
    who = join([f"{_initials(g)} {f}".strip() for g, f in people])          # ieee
    return num + f"{who}, \"{title},\" " + (f"{venue}, " if venue else "") + f"{year}." \
        + (f" [Online]. Available: {url}" if url else "")


def cited_order(ms: dict, spec: dict, n_papers: int) -> list[int]:
    seen = []
    for s in spec["sections"]:
        for m in CITE_RE.finditer((ms.get("sections") or {}).get(s["key"], "") or ""):
            for n in cite_numbers(m.group(1)):
                if 1 <= n <= n_papers and n not in seen:
                    seen.append(n)
    return seen


_OWN_LABEL = re.compile(r"^(?:(?:[A-Z]|[IVXLC]+|\d+)[.)]|\d+(?:\.\d+)+\.?)\s+(?=\S)")
_ACK = re.compile(r"acknowledge?ments?\b", re.I)
_UNNUMBERED = re.compile(r"appendi(x|ces)\b|supplementa", re.I)


def unlabel(text: str) -> str:
    """Drop a number the writer typed on a subsection ("A. Method", "2.1 Method");
    the layout adds its own."""
    return _OWN_LABEL.sub("", text.strip(), count=1) or text


def statements(spec: dict, details: dict) -> list[tuple[str, str]]:
    """Back-matter statements: the ones the venue requires plus any the
    author filled in. Unfilled required ones become a [TO ADD] note."""
    filled = {
        "Acknowledgements": details.get("acknowledgements"),
        "Funding": details.get("funding"),
        "Conflict of interest": details.get("conflicts"),
        "Data availability": details.get("data_availability"),
    }
    patterns = {"Acknowledgements": r"acknowledg", "Funding": r"fund", "Conflict of interest": r"conflict|competing",
                "Data availability": r"data"}
    out, used = [], set()
    for req in spec.get("required_statements") or []:
        text = ""
        for name, pat in patterns.items():
            if re.search(pat, req, re.I):
                text, _ = (filled.get(name) or "").strip(), used.add(name)
                break
        if re.search(r"\bAI\b|generative|artificial intelligence", req, re.I):
            text, _ = AI_DISCLOSURE, used.add("ai")
        out.append((req, text or f"[TO ADD: {req}]"))
    for name, text in filled.items():
        if name not in used and (text or "").strip():
            out.append((name, text.strip()))
    if details.get("ai_disclosure", True) and "ai" not in used:
        out.append(("Declaration of generative AI use", AI_DISCLOSURE))
    return out


def _results_slot(spec: dict) -> str:
    for s in spec["sections"]:
        if re.search(r"result|finding|evaluation|experiment", s["heading"], re.I):
            return s["key"]
    return spec["sections"][-1]["key"]


def order_figures(spec: dict, uploaded: list[dict], drawn: list[dict]) -> list[dict]:
    """All figures in the order they appear in the paper, numbered 1..n.

    Drawn concept diagrams sit in the section they illustrate; uploaded
    figures sit in the results section. Each keeps `old_n` -- the number the
    writer used for it in the text ("Fig. 2") before diagrams were added --
    so the renderers can correct those references (see refigure)."""
    keys = [s["key"] for s in spec["sections"]]
    slot = _results_slot(spec)
    out = []
    for k in keys:
        for f in drawn:
            if (f.get("section") if f.get("section") in keys else keys[0]) == k:
                out.append({**f, "section": k, "old_n": None})
        if k == slot:
            out += [{**f, "section": k, "old_n": f["n"]} for f in uploaded]
    for n, f in enumerate(out, 1):
        f["n"] = n
        ext = f["filename"].rsplit(".", 1)[-1].lower()
        f["outname"] = f"fig{n}.{ext}"
    return out


def refigure(text: str, figures: list[dict]) -> str:
    """Rewrite the writer's "Fig. N" references to the final numbering."""
    remap = {f["old_n"]: f["n"] for f in figures if f.get("old_n") and f["old_n"] != f["n"]}
    if not remap:
        return text or ""
    return re.sub(r"\b(Fig(?:ure|\.)?\s*)(\d+)\b",
                  lambda m: m.group(1) + str(remap.get(int(m.group(2)), int(m.group(2)))), text or "")


# ── Checklist ───────────────────────────────────────────────────────────────

def _words(t: str) -> int:
    return len(re.findall(r"\b[\w'-]+\b", t or ""))


def check(spec: dict, ms: dict, papers: list[dict], details: dict, paper_type: str,
          n_results: int = 0) -> list[dict]:
    rows = []

    def add(label, ok, detail):
        rows.append({"label": label, "status": "ok" if ok else "warn", "detail": detail})

    secs = ms.get("sections") or {}
    aw, limit = _words(ms.get("abstract")), spec.get("abstract_word_limit")
    add("Abstract length", not limit or aw <= limit, f"{aw} words" + (f" of {limit} allowed" if limit else ""))
    total = sum(_words(secs.get(s["key"])) for s in spec["sections"])
    if spec.get("word_limit"):
        add("Word count", total <= spec["word_limit"], f"{total} words of {spec['word_limit']} allowed (body text)")
    else:
        add("Word count", True, f"{total} words in the body; the template states no limit")
    if spec.get("page_limit"):
        rows.append({"label": "Page limit", "status": "warn",
                     "detail": f"The template allows {spec['page_limit']} pages. Pages can only be counted "
                               "after the file is opened or compiled -- check it there."})
    fmts = spec.get("accepted_formats") or []
    if fmts and "word" not in fmts:
        names = {"latex": "LaTeX", "pdf": "PDF"}
        rows.append({"label": "Submission format", "status": "warn",
                     "detail": f"{spec.get('venue') or 'This venue'} accepts {' or '.join(names[f] for f in fmts)}, "
                               "not Word files. Use the LaTeX download, or save the Word file as PDF before submitting."})
    missing = [s["heading"] for s in spec["sections"] if _words(secs.get(s["key"])) < 20]
    add("Required sections", not missing,
        "All sections written" if not missing else "Empty or very short: " + ", ".join(missing))
    kw = ms.get("keywords") or []
    add("Keywords", 0 < len(kw) <= (spec.get("keywords_max") or 99), f"{len(kw)} keywords")

    cited = cited_order(ms, spec, len(papers))
    bad = sorted({n for s in spec["sections"] for m in CITE_RE.finditer(secs.get(s["key"], "") or "")
                  for n in cite_numbers(m.group(1)) if not 1 <= n <= len(papers)})
    add("Citations", not bad and bool(cited),
        (f"{len(cited)} of {len(papers)} sources cited; uncited sources are left out of the reference list"
         if not bad else f"Citation numbers with no matching source: {bad}"))

    todos = sum(len(TODO_RE.findall(secs.get(s['key'], '') or '')) for s in spec["sections"])
    todos += len(TODO_RE.findall(ms.get("abstract") or ""))
    add("Content still to add", todos == 0,
        "Nothing flagged" if todos == 0 else
        f"{todos} [TO ADD] note(s) mark where the supplied material did not cover what the section needs")
    if paper_type == "research":
        add("Results", n_results > 0,
            f"{n_results} results file(s) used" if n_results else
            "No results were supplied, so Results and Discussion contain [TO ADD] notes instead of findings")
    add("Authors", bool(parse_authors(details.get("authors_text"))),
        "Author details provided" if parse_authors(details.get("authors_text")) else "No authors entered")
    pend = [h for h, t in statements(spec, details) if t.startswith("[TO ADD")]
    if spec.get("required_statements"):
        add("Required statements", not pend, "All provided" if not pend else "Missing: " + ", ".join(pend))
    return rows


# ── Word ────────────────────────────────────────────────────────────────────

_FALLBACK_STYLES = {
    "title": ["Title", "paper title", "Paper Title"],
    "author": ["Author", "Authors"],
    "affiliation": ["Affiliation"],
    "abstract": ["Abstract"],
    "keywords": ["Keywords", "Key Words"],
    "heading1": ["Heading 1"],
    "heading2": ["Heading 2"],
    "body": ["Body Text", "Normal"],
    "reference": ["References", "Reference", "Bibliography"],
    "caption": ["Caption", "figure caption"],
    "table_caption": ["table head", "Table Caption", "Table Title", "table caption"],
}


def _roman(n: int) -> str:
    out = ""
    for v, r in ((10, "X"), (9, "IX"), (5, "V"), (4, "IV"), (1, "I")):
        while n >= v:
            out, n = out + r, n - v
    return out


def _exemplars(doc, learned_roles: dict | None = None) -> dict:
    """How the template itself formats each part of a paper, read from its
    own sample paragraphs.

    Many publisher templates (IEEE Transactions, for one) do not rely on
    named styles: their sample title, authors, abstract and body text are
    'Normal' paragraphs with the formatting applied directly. Writing into
    such a template by style name loses all of it. So for each role we find
    the template's sample paragraph and keep a copy of its paragraph
    properties and of the character formatting of its main run; new
    paragraphs are then built from those copies. Templates that do use named
    styles work the same way, since the style reference is part of what is
    copied."""
    import copy
    from collections import Counter
    from docx.oxml.ns import qn

    def pack(p, label=None):
        ppr = copy.deepcopy(p._p.pPr) if p._p.pPr is not None else None
        if ppr is not None:
            for tag in ("w:sectPr", "w:framePr"):
                el = ppr.find(qn(tag))
                if el is not None:
                    ppr.remove(el)
        runs = [r for r in p.runs if r.text.strip()]
        main = max(runs, key=lambda r: len(r.text)) if runs else None
        rpr = main._r.find(qn("w:rPr")) if main is not None else None
        first = runs[0]._r.find(qn("w:rPr")) if runs else None
        return {"pPr": ppr, "rPr": copy.deepcopy(rpr) if rpr is not None else None,
                "label_rPr": copy.deepcopy(first) if first is not None else None,
                "text": p.text.strip(), "label": label,
                "numbered": _auto_numbered(ppr)}

    def _auto_numbered(ppr):
        num = ppr.find(qn("w:numPr")) if ppr is not None else None
        nid = num.find(qn("w:numId")) if num is not None else None
        return nid is not None and nid.get(qn("w:val")) != "0"      # numId 0 switches numbering OFF

    paras = [p for p in doc.paragraphs if p.text.strip()]
    sname = lambda p: (p.style.name if p.style is not None else "").lower()  # noqa: E731
    ex: dict = {}
    if not paras:
        return ex

    def find(pred, start=0):
        return next((i for i in range(start, len(paras)) if pred(paras[i])), None)

    t = find(lambda p: "title" in sname(p) and "sub" not in sname(p))
    t = 0 if t is None or t > 3 else t
    ex["title"] = pack(paras[t])
    a = find(lambda p: re.match(r"abstract\s*[—:\-–.]", p.text.strip(), re.I) and len(p.text) > 60, t + 1)
    if a is not None:
        m = re.match(r"(abstract\s*[—:\-–.]\s?)", paras[a].text.strip(), re.I)
        ex["abstract"] = pack(paras[a], m.group(1))
        if a - t >= 2:
            ex["author"] = pack(paras[a - 1])        # the line just above the abstract; notes can sit under the title
    k = find(lambda p: re.match(r"(index terms|key\s?words?)\s*[—:\-–.]", p.text.strip(), re.I), t + 1)
    if k is not None:
        m = re.match(r"((index terms|key\s?words?)\s*[—:\-–.]\s?)", paras[k].text.strip(), re.I)
        ex["keywords"] = pack(paras[k], m.group(1))
    h1 = find(lambda p: re.fullmatch(r"heading\s*1", sname(p)) is not None)
    if h1 is not None:
        ex["heading1"] = pack(paras[h1])
    h2 = find(lambda p: re.fullmatch(r"heading\s*2", sname(p)) is not None)
    if h2 is not None:
        ex["heading2"] = pack(paras[h2])

    start = (h1 if h1 is not None else (a if a is not None else t)) + 1
    def sig(p):
        raw = p._p.pPr.xml if p._p.pPr is not None else ""
        return re.sub(r'\sw:rsid\w*="[^"]*"', "", raw)
    cands = [p for p in paras[start:] if len(p.text) > 180 and "heading" not in sname(p)
             and (p._p.pPr is None or p._p.pPr.find(qn("w:framePr")) is None)
             and not re.match(r"\[\d+\]", p.text.strip())]
    if cands:
        best = Counter(sig(p) for p in cands).most_common(1)[0][0]
        ex["body"] = pack(next(p for p in cands if sig(p) == best))

    rh = find(lambda p: re.fullmatch(r"(references?|acknowledge?ments?)", p.text.strip(), re.I) is not None, start)
    if rh is not None:
        ex["back_heading"] = pack(paras[rh])
    r1 = find(lambda p: re.match(r"\[1\]\s", p.text.strip()) is not None, start)
    if r1 is None:
        # IEEE Transactions: sample references are plain paragraphs numbered by a list ("[1]" is not typed)
        r1 = find(lambda p: _auto_numbered(p._p.pPr) and "heading" not in sname(p)
                  and re.search(r"[“\"].{8,}[”\"]", p.text) and re.search(r"\b(19|20)\d\d\b", p.text), start)
    if r1 is not None:
        ex["reference"] = pack(paras[r1])
    c = find(lambda p: re.match(r"fig(ure)?\.?\s*1\b", p.text.strip(), re.I) and len(p.text) < 400, start)
    if c is not None:
        ex["caption"] = pack(paras[c])

    # What the model learned by reading the template (TemplateFormatAgent)
    # takes precedence over the rules of thumb above, which stay as the
    # fallback for a role it did not identify. A light sanity check guards
    # against an obviously wrong pick.
    def label_of(text, role):
        m = re.match(r"([A-Za-z][A-Za-z ]{2,22}\s*[—:\-–.]\s?)", text) if role in ("abstract", "keywords") else None
        return m.group(1) if m and len(text) > len(m.group(1)) + 20 else None

    for role, i in (learned_roles or {}).items():
        if not isinstance(i, int) or not 0 <= i < len(paras):
            continue
        text = paras[i].text.strip()
        if role in ("heading1", "heading2", "back_heading", "title") and len(text) > 220:
            continue
        if role == "body" and len(text) < 60:
            continue
        ex[role] = pack(paras[i], label_of(text, role))
    return ex


def _typed_number(sample: str) -> str | None:
    """How a template TYPES its heading numbers, read from its sample
    heading: 'roman' (IV.), 'alpha' (B.), 'arabic' (2.), 'decimal' (2.1)."""
    s = sample.strip()
    if re.match(r"\d+\.\d+\.?\s", s):
        return "decimal"
    if re.match(r"\d+[.)]?\s+\S", s):
        return "arabic"
    if re.match(r"[IVXLC]+[.)]\s", s):
        return "roman"
    if re.match(r"[A-Z][.)]\s", s):
        return "alpha"
    return None


def render_docx(template: bytes, spec: dict, ms: dict, papers: list[dict], details: dict,
                figures: list[dict], tables: list[dict], info: dict) -> bytes:
    import docx
    from docx.oxml.ns import qn
    from docx.shared import Inches

    import copy

    doc = docx.Document(io.BytesIO(as_docx_bytes(template)))
    body = doc.element.body
    W_P, W_PPR, W_SECT = qn("w:p"), qn("w:pPr"), qn("w:sectPr")

    # built-in layout: its styles are the format, nothing to learn
    EX = {} if info.get("number_headings") else _exemplars(doc, (info.get("learned") or {}).get("roles"))

    # Which of the template's sections is which (see docx_sections): the one
    # that held the most sample text is the body; the first is the title
    # block; one before the body written in an author style is the author
    # block. Their page setup and columns are kept, every other section break
    # is dropped, and all sample text is removed.
    sects = docx_sections(doc)
    final = sects[-1]["marker"]
    body_i = max(range(len(sects)), key=lambda i: sects[i]["chars"])
    author_i = next((i for i in range(1, body_i) if sects[i]["author_chars"] > 0), 0)
    if body_i == 0:
        author_i = 0
    keep = {0, author_i, body_i}
    n_authors = len(parse_authors(details.get("authors_text")))
    if author_i not in (0, body_i) and n_authors:
        cols = sects[author_i]["sect"].find(qn("w:cols"))       # e.g. IEEE: 4 columns for up to 6 authors
        if cols is not None:
            for c in list(cols):
                cols.remove(c)
            cols.set(qn("w:num"), str(max(1, min(n_authors, sects[author_i]["cols"], 3))))
            cols.set(qn("w:equalWidth"), "1")
    if sects[body_i]["marker"] is not final:                    # make the body the document's last section
        src, dst = sects[body_i]["sect"], sects[-1]["sect"]
        for c in list(dst):
            dst.remove(c)
        for c in src:
            dst.append(copy.deepcopy(c))
        keep.discard(body_i)
    markers = {i: (sects[i]["marker"] if i in keep else None) for i in range(len(sects))}
    kept = {id(m) for m in markers.values() if m is not None}
    for el in list(body.iterchildren()):
        if el is final:
            continue
        if id(el) in kept:
            for child in list(el):                              # keep the break, drop its sample text
                if child.tag != W_PPR:
                    el.remove(child)
            continue
        body.remove(el)
    front = markers[0] if body_i != 0 else None
    auth_anchor = markers[author_i] if author_i not in (0, body_i) else front

    para_styles = {s.name for s in doc.styles if getattr(s, "type", None) == 1}
    chosen = spec.get("styles") or {}

    def style(role):
        for name in [chosen.get(role)] + _FALLBACK_STYLES.get(role, []):
            if name and name in para_styles:
                return name
        return None

    numeric = spec.get("citation_style") != "author-year"
    order = cited_order(ms, spec, len(papers))
    renum = {old: new for new, old in enumerate(order, 1)}

    def cites(text):
        def sub(m):
            nums = cite_numbers(m.group(1))
            if not nums or any(n not in renum for n in nums):
                return m.group(0)
            if numeric:
                return "[" + ", ".join(str(n) for n in sorted({renum[n] for n in nums})) + "]"
            return "(" + "; ".join(author_year(papers[n - 1]) for n in nums) + ")"
        return CITE_RE.sub(sub, text or "")

    _EX_FALLBACK = {"affiliation": "author", "table_caption": "caption"}

    def add(text, role, before=None, bold_prefix=None):
        ex = EX.get(role) or EX.get(_EX_FALLBACK.get(role, ""))
        p = doc.add_paragraph(style=None if ex else style(role))
        if ex:                                       # format it the way the template's own sample is
            if p._p.pPr is not None:
                p._p.remove(p._p.pPr)
            if ex["pPr"] is not None:
                p._p.insert(0, copy.deepcopy(ex["pPr"]))

        def run(t, rpr, bold=False):
            r = p.add_run(t)
            if rpr is not None:
                old = r._r.find(qn("w:rPr"))
                if old is not None:
                    r._r.remove(old)
                r._r.insert(0, copy.deepcopy(rpr))
            elif bold:
                r.bold = True

        if bold_prefix:
            run((ex or {}).get("label") or bold_prefix, ex["label_rPr"] if ex else None, bold=True)
        run(text, ex["rPr"] if ex else None)
        if before is not None:
            before.addprevious(p._p)
        return p

    def like(role, text):
        """Match the template's own casing for headings (IEEE types them in capitals)."""
        sample = re.sub(r"^([IVXLC]+|[A-Z]|\d+(\.\d+)*)[.)]?\s+", "", (EX.get(role) or {}).get("text", ""))
        return text.upper() if sample and sample.isupper() else text

    # Does the template type its heading numbers by hand ("I. INTRODUCTION", "2.1 ...")? If so, in which scheme.
    h1_typed = None if (EX.get("heading1") or {}).get("numbered") else _typed_number((EX.get("heading1") or {}).get("text", ""))
    h2_typed = None if (EX.get("heading2") or {}).get("numbered") else _typed_number((EX.get("heading2") or {}).get("text", ""))
    if h1_typed == "alpha":
        h1_typed = None                              # "A Survey of ..." is a word, not a number

    def h1_label(i, text):
        return {"roman": f"{_roman(i)}. ", "arabic": f"{i}. "}.get(h1_typed, "") + text

    def h2_label(i, j, text):
        return {"alpha": f"{chr(64 + min(j, 26))}. ", "decimal": f"{i}.{j} ", "arabic": f"{j}. "}.get(h2_typed, "") + text

    add(ms.get("title") or "Untitled", "title", before=front)
    authors = parse_authors(details.get("authors_text"))
    if authors and EX.get("author") and auth_anchor is front:
        # one line, as the template's sample author line: "A. Author, Affiliation, B. Author, ..."
        add(", ".join(a["name"] + (f" ({a['affiliation']})" if a["affiliation"] else "") for a in authors),
            "author", before=auth_anchor)
    elif authors and auth_anchor is not front:
        for a in authors:                       # the template's own author block: one paragraph per author
            p = add(a["name"], "author", before=auth_anchor)
            for extra in (a["affiliation"], a["email"]):
                if extra:
                    p.add_run().add_break()
                    p.add_run(extra)
    elif authors:
        add(", ".join(a["name"] for a in authors), "author", before=auth_anchor)
        seen = []
        for a in authors:
            if a["affiliation"] and a["affiliation"] not in seen:
                seen.append(a["affiliation"])
                add(a["affiliation"], "affiliation" if style("affiliation") else "author", before=auth_anchor)
        mails = [a["email"] for a in authors if a["email"]]
        if mails:
            add(", ".join(mails), "affiliation" if style("affiliation") else "author", before=auth_anchor)
    else:
        add("[TO ADD: authors and affiliations]", "author", before=auth_anchor)

    if info.get("abstract_inline") or (EX.get("abstract") or {}).get("label"):
        add(ms.get("abstract") or "", "abstract", bold_prefix="Abstract—")
    else:
        add("Abstract", "heading1" if not style("abstract") else "abstract").runs[0].bold = True
        add(ms.get("abstract") or "", "abstract" if style("abstract") else "body")
    if ms.get("keywords"):
        add(", ".join(ms["keywords"]), "keywords" if (style("keywords") or EX.get("keywords")) else "body",
            bold_prefix="Keywords—" if info.get("abstract_inline") else "Keywords: ")

    def self_numbering(role):
        """True when the template's style for this role numbers itself (IEEE's
        'figure caption' prints 'Fig. 1.' on its own) -- then we must not."""
        name = style(role)
        if not name:
            return False
        ppr = doc.styles[name].element.find(qn("w:pPr"))
        return ppr is not None and ppr.find(qn("w:numPr")) is not None

    def put_figures(key):
        two_col = (info.get("columns") or 1) >= 2
        for f in [x for x in figures if x.get("section") == key]:
            if f.get("mention"):
                add(f["mention"].replace("{fig}", f"Fig. {f['n']}"), "body")
            # drawn diagrams are designed at one-column width; do not blow them up to a full page
            width = Inches(3.2 if two_col else (4.2 if f.get("old_n") is None else 5.6))
            try:
                p = doc.add_paragraph()
                p.add_run().add_picture(io.BytesIO(f["bytes"]), width=width)
                p.alignment = 1
            except Exception:  # noqa: BLE001 -- unreadable image: say so in the file
                add(f"[TO ADD: figure file {f['filename']} could not be embedded]", "body")
            add(f["caption"] if self_numbering("caption") else f"Fig. {f['n']}. {f['caption']}", "caption")

    def put_floats():
        width = Inches(3.2 if (info.get("columns") or 1) >= 2 else 5.6)
        for t in tables:
            if style("table_caption"):
                add(t["caption"] if self_numbering("table_caption") else f"Table {t['n']}. {t['caption']}",
                    "table_caption")
            else:
                add(f"Table {t['n']}. {t['caption']}", "body")
            cols = t["columns"]
            tbl = doc.add_table(rows=1 + len(t["rows"]), cols=len(cols))
            tbl.autofit = False
            tbl.alignment = 1
            for col in tbl.columns:                  # the grid widths too -- some viewers go by these
                col.width = int(width / len(cols))
            for row in tbl.rows:                     # fit the text column, not the page
                for cell in row.cells:
                    cell.width = int(width / len(cols))
            # Thin rules drawn directly, so the table reads as one whether or
            # not the template defines a table style.
            from docx.oxml import OxmlElement
            borders = OxmlElement("w:tblBorders")
            for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
                b = OxmlElement(f"w:{edge}")
                b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "4")
                b.set(qn("w:space"), "0"); b.set(qn("w:color"), "auto")
                borders.append(b)
            tbl._tbl.tblPr.append(borders)
            for j, c in enumerate(cols):
                tbl.cell(0, j).text = str(c)
            for i, row in enumerate(t["rows"], 1):
                for j, v in enumerate(row):
                    tbl.cell(i, j).text = "" if v is None else str(v)

    slot = _results_slot(spec) if tables else None
    # Templates number their own headings through their styles; the built-in
    # standard layout has plain styles, so number here ("2. Methods", "2.1 ...").
    numbered = bool(info.get("number_headings"))
    i = 0
    for s in spec["sections"]:
        if _ACK.match(s["heading"]):
            continue                            # printed once, with the back-matter statements
        head = like("heading1", s["heading"])
        if _UNNUMBERED.match(s["heading"]):     # appendices carry no section number
            add(head, "heading1")
        else:
            i += 1
            add(f"{i}. {head}" if numbered else h1_label(i, head), "heading1")
        sub = 0
        for kind, text in blocks(strip_own_heading((ms.get("sections") or {}).get(s["key"], ""), s["heading"])):
            if kind == "h2":
                sub += 1
                text = unlabel(text)
                add(f"{i}.{sub} {text}" if numbered else h2_label(i, sub, text), "heading2")
            else:
                add(cites(refigure(text, figures)), "body")
        if s["key"] == slot:
            put_floats()
        put_figures(s["key"])

    # IEEE and others set "References"/"Acknowledgment" in their own unnumbered
    # heading style; use it when the template showed us one.
    back = info.get("back_heading") if info.get("back_heading") in para_styles else None

    def back_heading(text):
        if EX.get("back_heading"):
            return add(like("back_heading", text), "back_heading")
        if back:
            p = doc.add_paragraph(style=back)
            p.add_run(text)
            return p
        return add(like("heading1", text), "heading1")

    for heading, text in statements(spec, details):
        back_heading(heading)
        add(text, "body")

    back_heading("References")
    ref_prefix = not ((style("reference") and self_numbering("reference")) or (EX.get("reference") or {}).get("numbered"))
    rstyle = spec.get("reference_style") or ("ieee" if numeric else "apa")
    rrole = "reference" if (EX.get("reference") or style("reference")) else "body"
    if numeric:
        for new, old in enumerate(order, 1):
            add(format_reference(papers[old - 1], rstyle, new, numbered=ref_prefix), rrole)
    else:
        for old in sorted(order, key=lambda n: _surname(papers[n - 1].get("authors") or "")[0].lower()):
            add(format_reference(papers[old - 1], rstyle if rstyle in ("apa", "harvard") else "apa", None), rrole)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# ── LaTeX ───────────────────────────────────────────────────────────────────

_TEX_MAP = {"\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
            "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}"}


def tex(s) -> str:
    return "".join(_TEX_MAP.get(c, c) for c in str(s if s is not None else ""))


def bib_keys(papers: list[dict]) -> list[str]:
    keys, taken = [], {}
    for p in papers:
        base = re.sub(r"[^A-Za-z]", "", _surname(p.get("authors") or "")[0]) or "ref"
        base = f"{base.lower()}{p.get('year') or 'nd'}"
        taken[base] = taken.get(base, 0) + 1
        keys.append(base if taken[base] == 1 else f"{base}{chr(96 + taken[base])}")
    return keys


def _bibtex(papers: list[dict], keys: list[str], order: list[int]) -> str:
    out = []
    for n in order:
        p, k = papers[n - 1], keys[n - 1]
        names = [a.strip() for a in re.split(r",|;| and |&", p.get("authors") or "") if a.strip()]
        names = ["others" if re.fullmatch(r"et al\.?", a, re.I) else a for a in names] or ["Anonymous"]
        venue = p.get("venue") or ""
        kind = "article" if venue else "misc"
        fields = [("author", " and ".join(tex(a) for a in names)), ("title", "{" + tex(p.get("title") or "") + "}")]
        if venue:
            fields.append(("journal", tex(venue)))
        if p.get("year"):
            fields.append(("year", str(p["year"])))
        if p.get("url"):
            fields.append(("url", p["url"].replace("%", r"\%")))
        out.append(f"@{kind}{{{k},\n" + ",\n".join(f"  {a} = {{{b}}}" for a, b in fields) + "\n}\n")
    return "\n".join(out)


def _strip_cmd(src: str, cmd: str) -> str:
    """Remove every \\cmd[...]{...} (balanced braces) from src."""
    pat = re.compile(r"\\" + cmd + r"\s*(\[[^\]]*\]\s*)?\{")
    while True:
        m = pat.search(src)
        if not m:
            return src
        depth, i = 1, m.end()
        while i < len(src) and depth:
            c = src[i]
            if c == "\\":
                i += 2
                continue
            depth += (c == "{") - (c == "}")
            i += 1
        src = src[:m.start()] + src[i:]


_ABSTRACT_FIRST = ("acmart", "elsarticle", "revtex", "aastex", "achemso")
_BIBSTYLE_BY_CLASS = {"ieeetran": "IEEEtran", "acmart": "ACM-Reference-Format", "llncs": "splncs04",
                      "elsarticle": "elsarticle-num", "svjour3": "spmpsci", "sn-jnl": "sn-basic"}


def _fallback_head(main: str, info: dict, title: str, authors: list[dict], abstract: str,
                   keywords: list[str]) -> str:
    i = main.find("\\begin{document}")
    pre = main[:i] if i >= 0 else f"\\documentclass{{{info.get('documentclass') or 'article'}}}\n"
    for cmd in ("title", "author"):
        pre = _strip_cmd(pre, cmd)
    names = " \\and ".join(
        tex(a["name"]) + (f"\\\\ {tex(a['affiliation'])}" if a["affiliation"] else "")
        + (f"\\\\ {tex(a['email'])}" if a["email"] else "") for a in authors
    ) or "[TO ADD: authors]"
    pre = pre.rstrip() + f"\n\n\\title{{{tex(title)}}}\n\\author{{{names}}}\n\n"
    abs_env = f"\\begin{{abstract}}\n{tex(abstract)}\n\\end{{abstract}}\n"
    cls = (info.get("documentclass") or "").lower()
    kw = f"\n\\noindent\\textbf{{Keywords:}} {tex(', '.join(keywords))}\n" if keywords else ""
    if any(cls.startswith(c) for c in _ABSTRACT_FIRST):
        return pre + "\\begin{document}\n" + abs_env + "\\maketitle\n" + kw
    return pre + "\\begin{document}\n\\maketitle\n" + abs_env + kw


def latex_head_for_agent(files: dict[str, bytes], info: dict) -> str | None:
    """The template's opening (to the first \\section) for LatexFrontMatterAgent,
    or None when there is nothing usable / it is too long to rewrite safely."""
    main = info.get("main_file")
    if not main or main not in files:
        return None
    src = _decode(files[main])
    m = re.search(r"^[^%\n]*\\section\*?\s*\{", src, re.M)
    end = m.start() if m else src.find("\\end{document}")
    head = src[:end] if end and end > 0 else src
    return head if "\\begin{document}" in head and len(head) <= 14000 else None


def _valid_head(head: str | None) -> bool:
    return bool(head and "\\documentclass" in head and "\\begin{document}" in head
                and "%%ABSTRACT%%" in head and "\\end{document}" not in head
                and not re.search(r"\\section\*?\s*\{", head))


def render_latex(template_name: str, template: bytes, info: dict, spec: dict, ms: dict,
                 papers: list[dict], details: dict, figures: list[dict], tables: list[dict],
                 agent_head: str | None = None) -> tuple[bytes, bool]:
    """Returns (zip bytes, used_template_front_matter)."""
    files = unpack_latex(template_name, template)
    main = info.get("main_file")
    main_src = _decode(files[main]) if main and main in files else ""
    authors = parse_authors(details.get("authors_text"))
    keys = bib_keys(papers)
    order = cited_order(ms, spec, len(papers))
    valid = set(order)

    def body_tex(text):
        out, pos = [], 0
        for m in CITE_RE.finditer(text or ""):
            nums = cite_numbers(m.group(1))
            out.append(tex(text[pos:m.start()]))
            out.append("\\cite{" + ",".join(keys[n - 1] for n in nums) + "}"
                       if nums and all(n in valid for n in nums) else tex(m.group(0)))
            pos = m.end()
        out.append(tex((text or "")[pos:]))
        return re.sub(r"(\[TO ADD:[^\]]*\])", r"\\textbf{\1}", "".join(out))

    used_agent = _valid_head(agent_head)
    head = (agent_head.replace("%%ABSTRACT%%", tex(ms.get("abstract") or "")) if used_agent else
            _fallback_head(main_src, info, ms.get("title") or "Untitled", authors,
                           ms.get("abstract") or "", ms.get("keywords") or []))
    if (figures or tables) and "graphicx" not in head:
        head = head.replace("\\begin{document}", "\\usepackage{graphicx}\n\\begin{document}", 1)
    if not re.search(r"\\usepackage(\[[^\]]*\])?\{[^}]*(url|hyperref)[^}]*\}", head):
        # reference URLs contain _ and %, which only typeset inside \url
        head = head.replace("\\begin{document}", "\\usepackage{url}\n\\begin{document}", 1)

    def floats():
        out = []
        for t in tables:
            cols = t["columns"]
            tab = ["\\begin{tabular}{" + "l" * len(cols) + "}", "\\hline",
                   " & ".join(tex(c) for c in cols) + " \\\\", "\\hline"]
            tab += [" & ".join(tex("" if v is None else v) for v in row) + " \\\\" for row in t["rows"]]
            tab += ["\\hline", "\\end{tabular}"]
            inner = "\n".join(tab)
            if len(cols) > 5:
                inner = "\\resizebox{\\linewidth}{!}{%\n" + inner + "}"
            out.append(f"\\begin{{table}}[t]\n\\centering\n\\caption{{{tex(t['caption'])}}}\\label{{tab:{t['n']}}}\n"
                       f"{inner}\n\\end{{table}}\n")
        return out

    def figs(key):
        out = []
        for f in [x for x in figures if x.get("section") == key]:
            if f.get("mention"):
                out += [tex(f["mention"]).replace("\\{fig\\}", f"Fig.~\\ref{{fig:{f['n']}}}"), ""]
            w = "0.62\\linewidth" if (f.get("old_n") is None and not two_col) else "\\linewidth"
            out.append(f"\\begin{{figure}}[t]\n\\centering\n\\includegraphics[width={w}]{{figures/{f['outname']}}}\n"
                       f"\\caption{{{tex(f['caption'])}}}\\label{{fig:{f['n']}}}\n\\end{{figure}}\n")
        return out

    two_col = bool(re.search(r"twocolumn|conference|journal|sigconf", info.get("class_options") or "", re.I)) \
        or (info.get("documentclass") or "").lower().startswith("ieeetran")
    parts = [head.rstrip(), ""]
    slot = _results_slot(spec) if tables else None
    for s in spec["sections"]:
        if _ACK.match(s["heading"]):
            continue
        parts.append(f"\\section{'*' if _UNNUMBERED.match(s['heading']) else ''}{{{tex(s['heading'])}}}")
        for kind, text in blocks(strip_own_heading((ms.get("sections") or {}).get(s["key"], ""), s["heading"])):
            parts.append(f"\\subsection{{{tex(unlabel(text))}}}" if kind == "h2" else body_tex(refigure(text, figures)))
            parts.append("")
        if s["key"] == slot:
            parts.extend(floats())
        parts.extend(figs(s["key"]))
    for heading, text in statements(spec, details):
        parts += [f"\\section*{{{tex(heading)}}}", body_tex(text), ""]

    cls = (info.get("documentclass") or "").lower()
    if info.get("biblatex"):
        if "\\addbibresource" in head:
            parts[0] = re.sub(r"\\addbibresource\{[^}]*\}", r"\\addbibresource{references.bib}", parts[0])
        else:
            parts[0] = parts[0].replace("\\begin{document}", "\\addbibresource{references.bib}\n\\begin{document}", 1)
        parts.append("\\printbibliography")
    else:
        by_ref = {"ieee": "IEEEtran", "apa": "apalike", "harvard": "apalike", "acm": "ACM-Reference-Format"}
        bst = info.get("bibstyle") or next((v for k, v in _BIBSTYLE_BY_CLASS.items() if cls.startswith(k)), None) \
            or ("apalike" if spec.get("citation_style") == "author-year" else None) \
            or (by_ref.get(spec.get("reference_style") or "") if main else None) or "unsrt"
        if "\\bibliographystyle" not in head:
            parts.append(f"\\bibliographystyle{{{bst}}}")
        parts.append("\\bibliography{references}")
    parts += ["", "\\end{document}", ""]

    readme = (
        "manuscript.tex   the generated paper -- compile this file\n"
        "references.bib   the cited sources\n"
        "figures/         the figures you uploaded\n"
        "Everything else is your template, unchanged.\n\n"
        "Compile with pdflatex + bibtex (or upload this folder to Overleaf and set\n"
        "manuscript.tex as the main document).\n"
        + ("" if used_agent else
           "\nNOTE: the title and author block were written with standard \\title/\\author\n"
           "commands. If your template uses its own author commands, move the details\n"
           "into them (compare with the template's sample file).\n")
        + "\nSearch for 'TO ADD' to find the places that still need your input.\n"
    )
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        reserved = {"manuscript.tex", "references.bib", "README_manuscript.txt"}
        for name, b in files.items():
            if name not in reserved and not name.startswith(("figures/fig", "_builtin")):
                z.writestr(name, b)
        z.writestr("manuscript.tex", "\n".join(parts))
        z.writestr("references.bib", _bibtex(papers, keys, order))
        z.writestr("README_manuscript.txt", readme)
        for f in figures:
            z.writestr(f"figures/{f['outname']}", f["bytes"])
    return buf.getvalue(), used_agent
