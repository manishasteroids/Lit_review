"""
Template reader for the Publish step.

Takes the file a user uploads as their target venue's template -- a Word
.docx/.dotx, or LaTeX as a .zip / .tex / .cls -- and returns a plain dict
describing what the template asks for: its section headings, the paragraph
styles it uses (Word) or its document class and bibliography style (LaTeX).
No LLM here: this is the factual read. agents/manuscript.py turns it into
the editable format spec, and core/manuscript_render.py writes into it.
"""
from __future__ import annotations

import io
import re
import zipfile

MAX_TEMPLATE_BYTES = 15 * 1024 * 1024
_MAX_ZIP_FILES = 400
_MAX_ZIP_UNPACKED = 60 * 1024 * 1024

_DOTX_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
_DOCX_CT = b"application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
_DOTM_CT = b"application/vnd.ms-word.template.macroEnabledTemplate.main+xml"


class TemplateError(ValueError):
    """Anything wrong with the uploaded template itself -> a 4xx, not a 500."""


def as_docx_bytes(data: bytes) -> bytes:
    """python-docx refuses .dotx (a different content type on the same file
    format). Rewrite that one content-type entry so a Word *template* opens
    like a document."""
    try:
        zin = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise TemplateError("This is not a valid Word file.") from e
    names = zin.namelist()
    ct = zin.read("[Content_Types].xml") if "[Content_Types].xml" in names else b""
    is_template = _DOTX_CT in ct or _DOTM_CT in ct
    # "Strict Open XML" (Word's other save format -- the IEEE conference
    # template ships this way) uses purl.oclc.org namespaces python-docx does
    # not know. The content is the same; map the namespaces to the usual ones.
    is_strict = b"purl.oclc.org/ooxml" in (zin.read("_rels/.rels") if "_rels/.rels" in names else b"")
    if not is_template and not is_strict:
        return data
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            buf = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                buf = buf.replace(_DOTX_CT, _DOCX_CT).replace(_DOTM_CT, _DOCX_CT)
            if is_strict and item.filename.lower().endswith((".xml", ".rels")):
                buf = _strict_to_transitional(buf)
            zout.writestr(item, buf)
    return out.getvalue()


_STRICT_NS = [
    (rb"http://purl\.oclc\.org/ooxml/officeDocument/relationships",
     rb"http://schemas.openxmlformats.org/officeDocument/2006/relationships"),
    (rb"http://purl\.oclc\.org/ooxml/(officeDocument|wordprocessingml|drawingml|spreadsheetml|presentationml)/",
     rb"http://schemas.openxmlformats.org/\1/2006/"),
    (rb'\s+w:conformance="strict"', rb""),
]


def _strict_to_transitional(buf: bytes) -> bytes:
    for pat, sub in _STRICT_NS:
        buf = re.sub(pat, sub, buf)
    return buf


# ── Word ────────────────────────────────────────────────────────────────────

_ROLE_HINTS = {
    "title": ["paper title", "papertitle", "article title", "title"],
    "author": ["author", "authors"],
    "affiliation": ["affiliation", "address", "institution"],
    "abstract": ["abstract"],
    "keywords": ["keywords", "key words", "index terms"],
    "heading1": ["heading 1", "heading1", "head1", "section"],
    "heading2": ["heading 2", "heading2", "head2", "subsection"],
    "body": ["body text", "bodytext", "text", "para", "normal"],
    "reference": ["references", "reference", "bibliography", "refs"],
    "caption": ["figure caption", "caption"],
}


def _guess_styles(names: list[str], used: dict[str, int]) -> dict[str, str]:
    low = {n.lower(): n for n in names}
    out: dict[str, str] = {}
    for role, hints in _ROLE_HINTS.items():
        for h in hints:
            if h in low:                                   # exact name first
                out[role] = low[h]
                break
        else:
            cands = [n for n in names if any(h in n.lower() for h in hints)]
            if role == "title":
                cands = [n for n in cands if "sub" not in n.lower() and "caption" not in n.lower()]
            if role == "body":
                cands = []                                 # decided by usage below
            if cands:
                out[role] = max(cands, key=lambda n: (used.get(n, 0) > 0, -len(n)))
    if "body" not in out or used.get(out["body"], 0) == 0:
        taken = {v for k, v in out.items() if k != "body"}
        free = {n: c for n, c in used.items() if n not in taken}
        if free:
            out["body"] = max(free, key=free.get)          # the style carrying the most text
    return out


def docx_sections(doc) -> list[dict]:
    """The template's sections in order. A section's properties sit on the
    marker that ENDS it: a paragraph carrying <w:sectPr>, or the body-level
    <w:sectPr> for the last one. For each: how much sample text it held and
    how much of that was in an author style -- which is how the renderer
    tells the title block, the author block and the body apart."""
    from docx.oxml.ns import qn
    W_P, W_PPR, W_SECT, W_T = qn("w:p"), qn("w:pPr"), qn("w:sectPr"), qn("w:t")
    styles = {s.style_id: (s.name or "") for s in doc.styles if getattr(s, "type", None) == 1}
    out, chars, author = [], 0, 0

    def close(marker, sect):
        nonlocal chars, author
        cols = sect.find(qn("w:cols"))
        num = cols.get(qn("w:num")) if cols is not None else None
        out.append({"marker": marker, "sect": sect, "chars": chars, "author_chars": author,
                    "cols": int(num) if (num or "").isdigit() else 1})
        chars, author = 0, 0

    for el in doc.element.body.iterchildren():
        if el.tag == W_SECT:
            close(el, el)
            continue
        n = sum(len(t.text or "") for t in el.iter(W_T))
        chars += n
        if el.tag == W_P:
            ppr = el.find(W_PPR)
            pstyle = ppr.find(qn("w:pStyle")) if ppr is not None else None
            if pstyle is not None and "author" in styles.get(pstyle.get(qn("w:val")), "").lower():
                author += n
            if ppr is not None and ppr.find(W_SECT) is not None:
                close(el, ppr.find(W_SECT))
    return out


def _page_furniture(doc) -> list[str]:
    """Text the template keeps in page headers/footers (e.g. IEEE's
    'REPLACE THIS LINE WITH YOUR MANUSCRIPT ID NUMBER'), which stays in the
    generated file and the author usually has to change."""
    seen = []
    for sec in doc.sections:
        for part in (sec.header, sec.first_page_header, sec.footer, sec.first_page_footer):
            try:
                text = " ".join(p.text.strip() for p in part.paragraphs if p.text.strip())
            except Exception:  # noqa: BLE001
                text = ""
            if len(text) > 8 and text not in seen:
                seen.append(text[:200])
    return seen[:6]


def content_paragraphs(doc) -> list:
    """The paragraphs every 'paragraph number' in this feature refers to:
    body paragraphs with text, in order."""
    return [p for p in doc.paragraphs if p.text.strip()]


def _effective(p) -> str:
    """Alignment, size and emphasis as they actually apply to a paragraph --
    direct formatting first, then its style chain."""
    def chain(style):
        while style is not None:
            yield style
            style = style.base_style
    from docx.oxml.ns import qn as _qn

    def jc(ppr):
        # read the raw value: Strict-format files use "start"/"end", which python-docx cannot map
        el = ppr.find(_qn("w:jc")) if ppr is not None else None
        return el.get(_qn("w:val")) if el is not None else None
    align = jc(p._p.pPr) or next((v for v in (jc(st.element.find(_qn("w:pPr"))) for st in chain(p.style)) if v), None)
    align = {"start": "left", "end": "right", "both": "justified", "distribute": "justified"}.get(align, align)
    runs = [r for r in p.runs if r.text.strip()]
    main = max(runs, key=lambda r: len(r.text)) if runs else None

    def font(attr):
        v = getattr(main.font, attr) if main is not None else None
        if v is None:
            v = next((getattr(s.font, attr) for s in chain(p.style) if getattr(s.font, attr) is not None), None)
        return v
    bits = [align or "left"]
    size = font("size")
    if size:
        bits.append(f"{size.pt:g}pt")
    for attr in ("bold", "italic", "all_caps", "small_caps"):
        if font(attr):
            bits.append(attr.replace("_", "-"))
    ppr = p._p.pPr
    from docx.oxml.ns import qn
    num = ppr.find(qn("w:numPr")) if ppr is not None else None
    nid = num.find(qn("w:numId")) if num is not None else None
    if nid is not None and nid.get(qn("w:val")) != "0":
        bits.append("auto-numbered")
    return ", ".join(bits)


def format_outline(doc, head: int = 115, tail: int = 70) -> list[str]:
    """The document as lines a model can read to learn its format:
    '#<paragraph number> [<style>] {<effective formatting>} <text>'.
    Long documents keep their start and their end (where references and
    biographies live); the numbers stay the true paragraph numbers."""
    paras = content_paragraphs(doc)
    pick = range(len(paras)) if len(paras) <= head + tail else \
        list(range(head)) + list(range(len(paras) - tail, len(paras)))
    out, last = [], -1
    for i in pick:
        if i != last + 1:
            out.append("... (paragraphs omitted) ...")
        p = paras[i]
        sname = p.style.name if p.style is not None else "Normal"
        out.append(f"#{i} [{sname}] {{{_effective(p)}}} {p.text.strip()[:110]}")
        last = i
    return out


def read_docx(data: bytes) -> dict:
    import docx
    from docx.oxml.ns import qn

    try:
        doc = docx.Document(io.BytesIO(as_docx_bytes(data)))
    except TemplateError:
        raise
    except Exception as e:  # noqa: BLE001
        raise TemplateError(f"Could not open this Word template: {e}") from e

    names = [s.name for s in doc.styles if getattr(s, "type", None) == 1 and s.name]
    used: dict[str, int] = {}
    outline, headings = [], []
    back_heading = None        # the style the template gives "References" / "Acknowledgment"
    for p in doc.paragraphs:
        text = (p.text or "").strip()
        if not text:
            continue
        sname = p.style.name if p.style is not None else "Normal"
        used[sname] = used.get(sname, 0) + len(text)
        if len(outline) < 140:
            outline.append({"style": sname, "text": text[:160]})
        if len(text) < 40 and re.match(r"(references?|acknowledge?ments?)\b", text, re.I):
            back_heading = sname
        if re.match(r"heading\s*1$|.*\bhead(ing)?\s*1\b", sname, re.I) and len(text) < 90:
            headings.append(text)

    sects = docx_sections(doc)                      # columns of the section that holds the body text
    columns = max(sects, key=lambda s: s["chars"])["cols"] if sects else 1

    styles = _guess_styles(names, used)
    abstract_inline = any(
        re.match(r"abstract\s*[—:\-–.]", o["text"], re.I) for o in outline
    )
    return {
        "kind": "docx",
        "style_names": names[:150],
        "styles": styles,
        "headings": headings[:30],
        "outline": outline,
        "columns": columns,
        "abstract_inline": abstract_inline,
        "back_heading": back_heading,
        "format_outline": format_outline(doc),
        "n_paragraphs": len(content_paragraphs(doc)),
        "page_furniture": _page_furniture(doc),
    }


# ── LaTeX ───────────────────────────────────────────────────────────────────

_TEXT_EXT = (".tex", ".cls", ".sty", ".bst", ".bib", ".bbx", ".cbx", ".clo", ".def", ".cfg", ".txt", ".md")


def unpack_latex(filename: str, data: bytes) -> dict[str, bytes]:
    """{relative path: bytes} for a .zip, or the single file for .tex/.cls."""
    lower = (filename or "").lower()
    if not lower.endswith(".zip"):
        return {re.sub(r"[^\w.\-]", "_", filename or "template.tex"): data}
    try:
        z = zipfile.ZipFile(io.BytesIO(data))
    except zipfile.BadZipFile as e:
        raise TemplateError("This .zip could not be opened.") from e
    files, total = {}, 0
    for item in z.infolist():
        name = item.filename.replace("\\", "/")
        if item.is_dir() or name.startswith("/") or ".." in name.split("/"):
            continue
        if name.startswith("__MACOSX/") or name.split("/")[-1].startswith("."):
            continue
        total += item.file_size
        if len(files) >= _MAX_ZIP_FILES or total > _MAX_ZIP_UNPACKED:
            raise TemplateError("This template archive is too large to use.")
        files[name] = z.read(item)
    if not files:
        raise TemplateError("This .zip is empty.")
    # A single top-level folder is just packaging -- drop it.
    tops = {n.split("/")[0] for n in files}
    if len(tops) == 1 and all("/" in n for n in files):
        files = {n.split("/", 1)[1]: b for n, b in files.items()}
    return files


def _decode(b: bytes) -> str:
    try:
        return b.decode("utf-8")
    except UnicodeDecodeError:
        return b.decode("latin-1", errors="replace")


def strip_comments(tex: str) -> str:
    return re.sub(r"(?<!\\)%.*", "", tex)


def find_main_tex(files: dict[str, bytes]) -> str | None:
    best, best_score = None, -1
    for name, b in files.items():
        if not name.lower().endswith(".tex"):
            continue
        t = strip_comments(_decode(b))
        if "\\documentclass" not in t or "\\begin{document}" not in t:
            continue
        base = name.split("/")[-1].lower()
        score = 10 - name.count("/")
        if re.search(r"main|template|sample|paper|manuscript|conference|article", base):
            score += 3
        if score > best_score:
            best, best_score = name, score
    return best


def read_latex(filename: str, data: bytes) -> dict:
    files = unpack_latex(filename, data)
    main = find_main_tex(files)
    cls_files = [n for n in files if n.lower().endswith(".cls")]
    info = {
        "kind": "latex",
        "files": sorted(files)[:200],
        "main_file": main,
        "documentclass": None,
        "class_options": "",
        "bibstyle": None,
        "biblatex": False,
        "headings": [],
        "main_excerpt": "",
    }
    if main:
        raw = _decode(files[main])
        t = strip_comments(raw)
        m = re.search(r"\\documentclass\s*(?:\[([^\]]*)\])?\s*\{([^}]+)\}", t)
        if m:
            info["class_options"], info["documentclass"] = (m.group(1) or "").strip(), m.group(2).strip()
        m = re.search(r"\\bibliographystyle\s*\{([^}]+)\}", t)
        if m:
            info["bibstyle"] = m.group(1).strip()
        info["biblatex"] = bool(re.search(r"\\usepackage(\[[^\]]*\])?\{biblatex\}|\\addbibresource", t))
        info["headings"] = [h.strip() for h in re.findall(r"\\section\*?\s*\{([^}]*)\}", t)][:30]
        info["main_excerpt"] = raw[:9000]
    elif cls_files:
        info["documentclass"] = cls_files[0].split("/")[-1][:-4]
    else:
        raise TemplateError(
            "No LaTeX template found in this upload -- it needs a .tex file with "
            "\\documentclass and \\begin{document}, or a .cls file."
        )
    return info


def read_template(filename: str, data: bytes) -> dict:
    if len(data) > MAX_TEMPLATE_BYTES:
        raise TemplateError(f"Template is over the {MAX_TEMPLATE_BYTES // (1024 * 1024)}MB limit.")
    lower = (filename or "").lower()
    if lower.endswith((".docx", ".dotx", ".dotm")):
        return read_docx(data)
    if lower.endswith((".zip", ".tex", ".cls")):
        return read_latex(filename, data)
    if lower.endswith(".doc"):
        raise TemplateError("Old .doc files are not supported -- save the template as .docx and upload that.")
    raise TemplateError("Upload a Word template (.docx or .dotx) or a LaTeX template (.zip, .tex or .cls).")


def describe(info: dict) -> str:
    """The template as text, for the format-spec agent's prompt."""
    if info.get("kind") == "docx":
        lines = [f"WORD TEMPLATE. Columns: {info.get('columns', 1)}.",
                 "Paragraphs in order, as [style] text:"]
        lines += [f"[{o['style']}] {o['text']}" for o in info.get("outline", [])]
        return "\n".join(lines)[:14000]
    return (
        f"LATEX TEMPLATE. Document class: {info.get('documentclass')} "
        f"[{info.get('class_options')}]. Bibliography style: {info.get('bibstyle')}. "
        f"Uses biblatex: {info.get('biblatex')}.\nMain file ({info.get('main_file')}):\n"
        f"{info.get('main_excerpt', '')}"
    )[:14000]
